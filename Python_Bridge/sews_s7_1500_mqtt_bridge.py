#!/usr/bin/env python3
"""MQTT <-> S7 DB100 bridge for SEWS S7-1500 PLCSIM Advanced tests."""

from __future__ import annotations

import argparse
import csv
import ctypes
import json
import os
import signal
import struct
import sys
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import paho.mqtt.client as mqtt

try:
    import snap7
except Exception as exc:  # pragma: no cover - shown to the user in console
    print("ERROR: python-snap7 is not installed. Run Scripts\\01_INSTALL_PYTHON_REQUIREMENTS.bat")
    raise

try:
    from openpyxl import Workbook, load_workbook
except Exception:
    Workbook = None
    load_workbook = None

try:
    import pyodbc
except Exception:
    pyodbc = None


CMD_INDEX = {
    "enable": 0,
    "auto_mode": 1,
    "enable_datalog": 2,
    "reset": 3,
    "force_maintenance_fault": 4,
    "force_machine_nok": 5,
    "electric_test_nok": 5,
    "adapter_test_nok": 5,
    "force_estop": 6,
    "scan_start": 7,
    "st2_done": 8,
    "st3_done": 9,
    "scan_finish": 10,
    "takeoff": 11,
    "harness_present": 12,
    "machine_ok": 13,
    "electric_test_ok": 13,
    "adapter_test_ok": 13,
    "bridge_active": 14,
    "transfer_to_electrical": 15,
    "move_to_electrical": 15,
    "electrical_harness_present": 16,
    "electrical_table_present": 16,
    "plug_connectors": 17,
    "connect_male_ports": 17,
    "electrical_ports_plugged": 17,
    "clip_finish_scan": 18,
    "electric_start_scan": 19,
    "electric_finish_scan": 20,
    "prebloc_cable_present": 21,
    "prebloc_start": 21,
    "prebloc_monter_cable": 21,
    "prebloc_done": 22,
    "prebloc_to_chariot": 22,
    "prebloc_dans_chariot": 22,
    "move_to_conveyor": 23,
    "chariot_to_conveyor": 23,
    "chariot_vers_convoyeur": 23,
    "conveyor_done": 24,
    "convoyeur_cablage_ok": 24,
    "audit_start": 25,
    "audit_ok": 26,
    "audit_nok": 27,
    "audit2_ok": 28,
    "audit3_ok": 29,
    "pack_ok": 30,
}

SIM_PANEL_DB_NUMBER = 1
SIM_PANEL_DB_SIZE = 36

PULSE_COMMANDS = {"reset", "scan_start", "st2_done", "st3_done", "transfer_to_electrical", "move_to_electrical", "plug_connectors", "connect_male_ports", "electrical_ports_plugged", "clip_finish_scan", "electric_start_scan", "electric_finish_scan", "scan_finish", "takeoff", "prebloc_cable_present", "prebloc_start", "prebloc_monter_cable", "prebloc_done", "prebloc_to_chariot", "prebloc_dans_chariot", "move_to_conveyor", "chariot_to_conveyor", "chariot_vers_convoyeur", "conveyor_done", "convoyeur_cablage_ok", "audit_start", "audit_ok", "audit_nok", "audit2_ok", "audit3_ok", "pack_ok"}
DEFAULT_PULSE_MS = {
    "reset": 500,
    "scan_start": 10000,
    "st2_done": 10000,
    "st3_done": 10000,
    "transfer_to_electrical": 10000,
    "move_to_electrical": 10000,
    "plug_connectors": 10000,
    "connect_male_ports": 10000,
    "electrical_ports_plugged": 10000,
    "clip_finish_scan": 10000,
    "electric_start_scan": 10000,
    "electric_finish_scan": 10000,
    "scan_finish": 10000,
    "takeoff": 10000,
    "prebloc_cable_present": 10000,
    "prebloc_start": 10000,
    "prebloc_monter_cable": 10000,
    "prebloc_done": 10000,
    "prebloc_to_chariot": 10000,
    "prebloc_dans_chariot": 10000,
    "move_to_conveyor": 10000,
    "chariot_to_conveyor": 10000,
    "chariot_vers_convoyeur": 10000,
    "conveyor_done": 10000,
    "convoyeur_cablage_ok": 10000,
    "audit_start": 10000,
    "audit_ok": 10000,
    "audit_nok": 10000,
    "audit2_ok": 10000,
    "audit3_ok": 10000,
    "pack_ok": 10000,
}

COMMAND_ALIASES = {
    "electric_test_ok": "machine_ok",
    "adapter_test_ok": "machine_ok",
    "electric_test_nok": "force_machine_nok",
    "adapter_test_nok": "force_machine_nok",
    "move_to_electrical": "transfer_to_electrical",
    "electrical_table_present": "electrical_harness_present",
    "connect_male_ports": "plug_connectors",
    "electrical_ports_plugged": "plug_connectors",
    "scan_clip_start": "scan_start",
    "scan_clips_start": "scan_start",
    "clip_start_scan": "scan_start",
    "scan_clip_finish": "clip_finish_scan",
    "scan_clips_finish": "clip_finish_scan",
    "scan_electric_start": "electric_start_scan",
    "scan_electrical_start": "electric_start_scan",
    "scan_electric_finish": "electric_finish_scan",
    "scan_electrical_finish": "electric_finish_scan",
    "prebloc_monter_cable": "prebloc_cable_present",
    "prebloc_start": "prebloc_cable_present",
    "godet_deposer_cable": "prebloc_cable_present",
    "deposer_dans_godet": "prebloc_cable_present",
    "godet_cable_present": "prebloc_cable_present",
    "prebloc_dans_chariot": "prebloc_done",
    "prebloc_to_chariot": "prebloc_done",
    "godet_cable_pris": "prebloc_done",
    "prendre_du_godet": "prebloc_done",
    "take_from_godet": "prebloc_done",
    "chariot_to_conveyor": "move_to_conveyor",
    "chariot_vers_convoyeur": "move_to_conveyor",
    "godet_to_conveyor": "move_to_conveyor",
    "convoyeur_cablage_ok": "conveyor_done",
    "audit_2_ok": "audit2_ok",
    "audit2": "audit2_ok",
    "audit_3_ok": "audit3_ok",
    "audit3": "audit3_ok",
    "packing_ok": "pack_ok",
    "pack_done": "pack_ok",
}

STATUS_BOOL_NAMES = [
    "harness_present",
    "machine_ok",
    "maintenance_alarm",
    "cycle_active",
    "fault_active",
    "st2_working",
    "st3_working",
    "st4_testing",
    "st5_takeoff_ready",
    "estop_active",
    "machine_nok",
    "trigger_scan_start",
    "trigger_scan_finish",
    "clip_table_harness_present",
    "electrical_table_harness_present",
    "electrical_ports_plugged",
    "transfer_to_electrical_ready",
    "electrical_current_passed",
    "trigger_scan_clip_start",
    "trigger_scan_clip_finish",
    "trigger_scan_electric_start",
    "trigger_scan_electric_finish",
    "prebloc_cable_present",
    "prebloc_ok",
    "chariot_loaded",
    "conveyor_table_present",
    "conveyor_ok",
    "audit_station_present",
    "audit_ok",
    "audit_nok",
    "audit_ready",
    "pack_ok",
]

STEP_TEXT = {
    0: "Ready - attente cable prebloc/godet",
    5: "Etape 1 - godet: cable prepare detecte",
    10: "Etape 1 OK - cable pris du godet",
    15: "Etape 2 - cable pose sur table convoyeur",
    18: "Etape 2 OK - pret table clips",
    20: "ST2 operator working",
    30: "ST3 operator working",
    34: "Clips done - waiting final clip scan",
    35: "Clip final scan OK - transfer to electrical table",
    37: "Electrical table - waiting electrical start scan",
    40: "Electrical test running - plug male ports into female adapters",
    50: "Electrical OK - waiting audit",
    55: "Etape 4 OK - waiting final audit",
    60: "Etape 5 - audit running",
    70: "Audit OK - takeoff ready",
    80: "Emergency stop",
    90: "Maintenance fault: sensor lost",
    91: "Electrical test NOK",
    92: "Audit NOK - quality control failed",
}

CYCLE_COLUMNS = [
    "cycle_no",
    "harness_code",
    "operator",
    "prebloc_start_time",
    "prebloc_done_time",
    "godet_loaded_time",
    "godet_sensor_detect_time",
    "godet_taken_time",
    "chariot_loaded_time",
    "prebloc_operation_sec",
    "prebloc_operation_min",
    "conveyor_start_time",
    "conveyor_done_time",
    "layup_start_time",
    "layup_done_time",
    "conveyor_operation_sec",
    "conveyor_operation_min",
    "deposit_time",
    "clip_table_time",
    "scan_start_time",
    "clip_scan_start_time",
    "tape_start_time",
    "tape_done_time",
    "tape_operation_sec",
    "tape_operation_min",
    "clip_first_clip_time",
    "st2_done_time",
    "st3_done_time",
    "clip_done_time",
    "clip_scan_finish_time",
    "clip_operation_sec",
    "clip_operation_min",
    "transfer_to_electrical_time",
    "electrical_table_time",
    "electric_scan_start_time",
    "connectors_plugged_time",
    "electric_work_start_time",
    "electric_test_time",
    "electric_scan_finish_time",
    "electric_operation_sec",
    "electric_operation_min",
    "audit_start_time",
    "audit_finish_time",
    "audit1_start_time",
    "audit1_finish_time",
    "audit2_finish_time",
    "audit3_finish_time",
    "pack_time",
    "pack_result",
    "audit_operation_sec",
    "audit_operation_min",
    "audit_result",
    "scan_finish_time",
    "takeoff_time",
    "reference_family",
    "required_clips_json",
    "clip_total_required",
    "clip_total_done",
    "clip_progress_text",
    "duration_work_sec",
    "duration_work_min",
    "plc_duration_sec",
    "fault_code",
    "fault_text",
    "fault_source",
    "result",
    "source",
    "sequence_errors",
    "comment",
]

REFERENCE_CLIP_MATRIX: dict[str, dict[str, int]] = {
    "24010-7TG0A": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 1, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TH0A": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF1E-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF2A-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF3A-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF3E-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF6A-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF6B-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF6C-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF6D-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF6E-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TF7A-09-0604": {"C1": 3, "C2": 4, "C3": 2, "C4": 3, "C5": 0, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 2, "C11": 1},
    "24010-7TH6A-09-0604": {"C1": 3, "C2": 4, "C3": 1, "C4": 3, "C5": 1, "C6": 9, "C7": 1, "C8": 5, "C9": 2, "C10": 1, "C11": 1},
}

CSV_DELIMITER = ";"
EXCEL_SEPARATOR_LINE = "sep=;"


def now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def parse_bool_payload(payload: Any, default: bool = True) -> bool:
    if isinstance(payload, dict):
        return parse_bool_payload(payload.get("value", default), default)
    if isinstance(payload, bool):
        return payload
    if isinstance(payload, (int, float)):
        return payload != 0
    if payload is None:
        return default
    text = str(payload).strip().lower()
    if text in {"0", "false", "off", "no", "nok"}:
        return False
    if text in {"1", "true", "on", "yes", "ok"}:
        return True
    return default


def normalize_reference(value: str) -> str:
    return "".join(ch.upper() if ch.isalnum() else "-" for ch in str(value or "").strip()).strip("-")


def reference_clip_plan(code: str) -> dict[str, int]:
    normalized = normalize_reference(code)
    if normalized in REFERENCE_CLIP_MATRIX:
        return dict(REFERENCE_CLIP_MATRIX[normalized])
    for reference, plan in REFERENCE_CLIP_MATRIX.items():
        if reference in normalized or normalized in reference:
            return dict(plan)
    return dict(next(iter(REFERENCE_CLIP_MATRIX.values())))


def reference_family(code: str) -> str:
    normalized = normalize_reference(code)
    for reference in REFERENCE_CLIP_MATRIX:
        if reference in normalized or normalized in reference:
            return reference
    return "SIMULATION-DEFAULT"


def clip_progress_text(done: dict[str, int], required: dict[str, int]) -> str:
    parts = []
    for clip in sorted(required, key=lambda value: int(value[1:])):
        target = int(required.get(clip, 0) or 0)
        current = int(done.get(clip, 0) or 0)
        if target:
            parts.append(f"{clip}:{current}/{target}")
    return " ".join(parts)


def set_s7_bool(buffer: bytearray, byte_index: int, bit_index: int, value: bool) -> None:
    mask = 1 << bit_index
    if value:
        buffer[byte_index] |= mask
    else:
        buffer[byte_index] &= (~mask) & 0xFF


def read_int(data: bytes, offset: int) -> int:
    return struct.unpack(">h", data[offset : offset + 2])[0]


def read_dint(data: bytes, offset: int) -> int:
    return struct.unpack(">i", data[offset : offset + 4])[0]


def read_real(data: bytes, offset: int) -> float:
    return struct.unpack(">f", data[offset : offset + 4])[0]


def pid_is_alive(pid: int) -> bool:
    if pid <= 0:
        return False
    handle = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)
    if handle:
        ctypes.windll.kernel32.CloseHandle(handle)
        return True
    return False


def acquire_named_mutex() -> int:
    if os.name != "nt":
        return 0
    handle = ctypes.windll.kernel32.CreateMutexW(None, False, "Local\\SEWS_S7_1500_MQTT_BRIDGE")
    if not handle:
        return 0
    if ctypes.windll.kernel32.GetLastError() == 183:
        ctypes.windll.kernel32.CloseHandle(handle)
        print("ERROR: another SEWS bridge instance is already running. Stop it first.", flush=True)
        return -1
    return int(handle)


def release_named_mutex(handle: int) -> None:
    if os.name == "nt" and handle > 0:
        try:
            ctypes.windll.kernel32.CloseHandle(handle)
        except Exception:
            pass


def acquire_lock(lock_path: Path) -> int:
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    if lock_path.exists():
        try:
            old_pid = int(lock_path.read_text(encoding="utf-8").strip() or "0")
        except Exception:
            old_pid = 0
        if pid_is_alive(old_pid):
            print(f"ERROR: another SEWS bridge is already running with PID {old_pid}. Stop it first.")
            return -1
        try:
            lock_path.unlink()
        except Exception:
            pass
    fd = os.open(str(lock_path), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    os.write(fd, str(os.getpid()).encode("ascii"))
    return fd


@dataclass
class BridgeConfig:
    config_path: Path
    s7_host: str
    s7_rack: int
    s7_slot: int
    db_number: int
    db_size: int
    poll_ms: int
    mqtt_host: str
    mqtt_port: int
    mqtt_client_id: str
    topic_prefix: str
    csv_path: Path
    xlsx_path: Path
    default_harness_code: str
    sql_enabled: bool
    sql_connection_string: str

    @classmethod
    def load(cls, path: str | os.PathLike[str]) -> "BridgeConfig":
        config_path = Path(path).resolve()
        raw = json.loads(config_path.read_text(encoding="utf-8"))
        base = config_path.parent
        data = raw.get("data", {})
        sql = raw.get("sql", {})
        csv_path = (base / data.get("csv_path", "cycles.csv")).resolve()
        xlsx_path = (base / data.get("xlsx_path", "cycles.xlsx")).resolve()
        s7 = raw.get("s7", {})
        mq = raw.get("mqtt", {})
        return cls(
            config_path=config_path,
            s7_host=s7.get("host", "127.0.0.1"),
            s7_rack=int(s7.get("rack", 0)),
            s7_slot=int(s7.get("slot", 1)),
            db_number=int(s7.get("db_number", 100)),
            db_size=int(s7.get("db_size", 320)),
            poll_ms=int(s7.get("poll_ms", 300)),
            mqtt_host=mq.get("host", "127.0.0.1"),
            mqtt_port=int(mq.get("port", 1883)),
            mqtt_client_id=mq.get("client_id", "s7_1500_demo_bridge"),
            topic_prefix=mq.get("topic_prefix", "s7_1500_demo").strip("/"),
            csv_path=csv_path,
            xlsx_path=xlsx_path,
            default_harness_code=data.get("default_harness_code", "DEMO-001"),
            sql_enabled=bool(sql.get("enabled", False)),
            sql_connection_string=sql.get(
                "connection_string",
                "DRIVER={ODBC Driver 18 for SQL Server};SERVER=localhost\\SQLEXPRESS;DATABASE=industrial_demo;Trusted_Connection=yes;TrustServerCertificate=yes;",
            ),
        )


@dataclass
class BridgeState:
    running: bool = True
    s7_connected: bool = False
    last_s7_error: str = ""
    current_code: str = ""
    last_status: dict[str, Any] = field(default_factory=dict)
    last_cycle_count: int = -1
    cycles: list[dict[str, Any]] = field(default_factory=list)
    current_cycle: dict[str, Any] = field(default_factory=dict)
    current_operator: str = ""
    current_comment: str = ""
    current_takt_target_sec: float = 60.0
    sequence_errors: int = 0
    pending_pulses: dict[int, float] = field(default_factory=dict)
    pending_web_commit_at: float = 0.0
    pending_web_cycle_marker: int = -1
    lock: threading.Lock = field(default_factory=threading.Lock)


class SewsBridge:
    def __init__(self, cfg: BridgeConfig) -> None:
        self.cfg = cfg
        self.state = BridgeState(current_code=cfg.default_harness_code)
        self.sql_available = False
        self.sql_last_error = ""
        self.sql_retry_at = 0.0
        self.sql_sync_at = 0.0
        self.last_sql_csv_mtime = 0.0
        self.client = snap7.client.Client()
        self.s7_lock = threading.RLock()
        self.mqtt = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2, client_id=cfg.mqtt_client_id)
        self.mqtt.on_connect = self.on_mqtt_connect
        self.mqtt.on_message = self.on_mqtt_message
        self.cfg.csv_path.parent.mkdir(parents=True, exist_ok=True)
        self.cfg.xlsx_path.parent.mkdir(parents=True, exist_ok=True)
        self.ensure_csv()
        self.state.cycles = self.load_cycles_from_csv()
        self.initialize_sql()
        if self.state.cycles:
            try:
                self.state.last_cycle_count = max(int(row.get("cycle_no") or 0) for row in self.state.cycles)
            except Exception:
                self.state.last_cycle_count = len(self.state.cycles)
        if os.environ.get("SEWS_SKIP_STARTUP_SQL_SYNC") != "1":
            self.sync_sql_from_csv("startup")

    def initialize_sql(self, quiet: bool = False) -> None:
        if not self.cfg.sql_enabled:
            return
        if pyodbc is None:
            self.sql_last_error = "pyodbc is not installed"
            if not quiet:
                print(f"SQL disabled: {self.sql_last_error}", flush=True)
            return
        try:
            with self.sql_connect() as conn:
                conn.cursor().execute("SELECT 1")
            self.sql_available = True
            self.sql_last_error = ""
            self.sql_retry_at = 0.0
            if not quiet:
                print("SQL: connected to SEWS database", flush=True)
        except Exception as exc:
            self.sql_available = False
            self.sql_last_error = str(exc)
            self.sql_retry_at = time.monotonic() + 5.0
            if not quiet:
                print(f"SQL disabled: {exc}", flush=True)

    def sql_connect(self):
        if pyodbc is None:
            raise RuntimeError("pyodbc is not installed")
        return pyodbc.connect(self.cfg.sql_connection_string, autocommit=True)

    def ensure_sql_available(self, force: bool = False) -> bool:
        if not self.cfg.sql_enabled:
            return False
        if self.sql_available:
            return True
        if not force and time.monotonic() < self.sql_retry_at:
            return False
        self.initialize_sql(quiet=not force)
        return self.sql_available

    def mark_sql_error(self, exc: Exception) -> None:
        self.sql_available = False
        self.sql_last_error = str(exc)
        self.sql_retry_at = time.monotonic() + 5.0
        print(f"SQL temporarily unavailable, retrying soon: {exc}", flush=True)

    @staticmethod
    def sql_float(value: Any) -> float | None:
        try:
            if value in ("", None):
                return None
            return float(value)
        except Exception:
            return None

    @staticmethod
    def sql_int(value: Any) -> int | None:
        try:
            if value in ("", None):
                return None
            return int(float(value))
        except Exception:
            return None

    @staticmethod
    def sql_bool(value: Any) -> int:
        return 1 if bool(value) else 0

    @staticmethod
    def sql_text(value: Any, limit: int = 500) -> str | None:
        if value in ("", None):
            return None
        return str(value)[:limit]

    @staticmethod
    def sql_dt(value: Any) -> str | None:
        if value in ("", None):
            return None
        text = str(value).strip()
        if not text:
            return None
        if text.endswith("Z"):
            text = text[:-1]
        return text.replace("T", " ")[:26]

    def sql_execute(self, query: str, params: tuple[Any, ...]) -> bool:
        if not self.ensure_sql_available():
            return False
        try:
            with self.sql_connect() as conn:
                conn.cursor().execute(query, params)
            self.sql_available = True
            self.sql_last_error = ""
            return True
        except Exception as exc:
            self.mark_sql_error(exc)
            return False

    def sql_scalar(self, query: str, params: tuple[Any, ...] = ()) -> Any:
        if not self.ensure_sql_available():
            return None
        try:
            with self.sql_connect() as conn:
                cursor = conn.cursor()
                row = cursor.execute(query, params).fetchone()
            self.sql_available = True
            self.sql_last_error = ""
            return row[0] if row else None
        except Exception as exc:
            self.mark_sql_error(exc)
            return None

    def sql_write_live_status(self, status: dict[str, Any]) -> None:
        if not self.cfg.sql_enabled or not self.sql_available:
            return
        query = """
        UPDATE dbo.Fact_LiveStatus SET
            ts = ?, connected = ?, step_no = ?, harness_code = ?,
            cycle_timer_sec = ?, last_duration_sec = ?,
            fault_code = ?, fault_source = ?, fault_text = ?,
            machine_nok = ?, maintenance_alarm = ?, updated_at = SYSDATETIME()
        WHERE id = 1;
        """
        self.sql_execute(
            query,
            (
                self.sql_dt(status.get("ts")),
                self.sql_bool(status.get("connected")),
                self.sql_int(status.get("step")),
                self.sql_text(status.get("harness_code"), 120),
                self.sql_float(status.get("cycle_timer_sec")),
                self.sql_float(status.get("last_duration_sec")),
                self.sql_int(status.get("fault_code")),
                self.sql_text(status.get("fault_source"), 80),
                self.sql_text(status.get("fault_text"), 500),
                self.sql_bool(status.get("machine_nok")),
                self.sql_bool(status.get("maintenance_alarm")),
            ),
        )

    def sql_write_cycle(self, row: dict[str, Any]) -> None:
        if not self.cfg.sql_enabled or not self.sql_available:
            return
        query = """
        MERGE dbo.Fact_Cycles WITH (HOLDLOCK) AS target
        USING (SELECT ? AS cycle_no) AS src
        ON target.cycle_no = src.cycle_no
        WHEN MATCHED THEN UPDATE SET
            harness_code = ?, operator_code = ?, reference_family = ?, result = ?, source = ?,
            fault_code = ?, fault_source = ?, fault_text = ?, sequence_errors = ?, comment = ?,
            prebloc_start_time = ?, prebloc_done_time = ?, conveyor_start_time = ?, conveyor_done_time = ?,
            tape_start_time = ?, tape_done_time = ?, clip_first_clip_time = ?, clip_scan_finish_time = ?,
            electric_work_start_time = ?, electric_scan_finish_time = ?, audit_start_time = ?, audit_finish_time = ?,
            takeoff_time = ?, prebloc_operation_sec = ?, conveyor_operation_sec = ?, tape_operation_sec = ?,
            clip_operation_sec = ?, electric_operation_sec = ?, audit_operation_sec = ?, duration_work_sec = ?,
            duration_work_min = ?, plc_duration_sec = ?, clip_total_required = ?, clip_total_done = ?,
            clip_progress_text = ?, updated_at = SYSDATETIME()
        WHEN NOT MATCHED THEN INSERT (
            cycle_no, harness_code, operator_code, reference_family, result, source,
            fault_code, fault_source, fault_text, sequence_errors, comment,
            prebloc_start_time, prebloc_done_time, conveyor_start_time, conveyor_done_time,
            tape_start_time, tape_done_time, clip_first_clip_time, clip_scan_finish_time,
            electric_work_start_time, electric_scan_finish_time, audit_start_time, audit_finish_time,
            takeoff_time, prebloc_operation_sec, conveyor_operation_sec, tape_operation_sec,
            clip_operation_sec, electric_operation_sec, audit_operation_sec, duration_work_sec,
            duration_work_min, plc_duration_sec, clip_total_required, clip_total_done, clip_progress_text
        ) VALUES (
            ?, ?, ?, ?, ?, ?,
            ?, ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?, ?
        );
        """
        values = (
            self.sql_int(row.get("cycle_no")),
            self.sql_text(row.get("harness_code"), 120),
            self.sql_text(row.get("operator"), 80),
            self.sql_text(row.get("reference_family"), 80),
            self.sql_text(row.get("result"), 30),
            self.sql_text(row.get("source"), 80),
            self.sql_int(row.get("fault_code")),
            self.sql_text(row.get("fault_source"), 80),
            self.sql_text(row.get("fault_text"), 500),
            self.sql_int(row.get("sequence_errors")),
            self.sql_text(row.get("comment"), 500),
            self.sql_dt(row.get("prebloc_start_time")),
            self.sql_dt(row.get("prebloc_done_time")),
            self.sql_dt(row.get("conveyor_start_time")),
            self.sql_dt(row.get("conveyor_done_time")),
            self.sql_dt(row.get("tape_start_time")),
            self.sql_dt(row.get("tape_done_time")),
            self.sql_dt(row.get("clip_first_clip_time")),
            self.sql_dt(row.get("clip_scan_finish_time")),
            self.sql_dt(row.get("electric_work_start_time")),
            self.sql_dt(row.get("electric_scan_finish_time")),
            self.sql_dt(row.get("audit_start_time")),
            self.sql_dt(row.get("audit_finish_time")),
            self.sql_dt(row.get("takeoff_time")),
            self.sql_float(row.get("prebloc_operation_sec")),
            self.sql_float(row.get("conveyor_operation_sec")),
            self.sql_float(row.get("tape_operation_sec")),
            self.sql_float(row.get("clip_operation_sec")),
            self.sql_float(row.get("electric_operation_sec")),
            self.sql_float(row.get("audit_operation_sec")),
            self.sql_float(row.get("duration_work_sec")),
            self.sql_float(row.get("duration_work_min")),
            self.sql_float(row.get("plc_duration_sec")),
            self.sql_int(row.get("clip_total_required")),
            self.sql_int(row.get("clip_total_done")),
            self.sql_text(row.get("clip_progress_text"), 500),
        )
        self.sql_execute(query, values[:1] + values[1:] + values)
        self.sql_write_operations(row)
        if self.sql_int(row.get("fault_code")):
            self.sql_write_maintenance(row)

    def sql_write_operations(self, row: dict[str, Any]) -> None:
        operations = [
            ("Preblock + godet", 1, "prebloc_start_time", "prebloc_done_time", "prebloc_operation_sec"),
            ("Cablage convoyeur", 2, "conveyor_start_time", "conveyor_done_time", "conveyor_operation_sec"),
            ("Tape", 3, "tape_start_time", "tape_done_time", "tape_operation_sec"),
            ("Table clips", 4, "clip_first_clip_time", "clip_scan_finish_time", "clip_operation_sec"),
            ("Test electrique", 5, "electric_work_start_time", "electric_scan_finish_time", "electric_operation_sec"),
            ("Audit final", 6, "audit_start_time", "audit_finish_time", "audit_operation_sec"),
        ]
        query = """
        MERGE dbo.Fact_Operations WITH (HOLDLOCK) AS target
        USING (SELECT ? AS cycle_no, ? AS operation_name) AS src
        ON target.cycle_no = src.cycle_no AND target.operation_name = src.operation_name
        WHEN MATCHED THEN UPDATE SET
            harness_code = ?, operator_code = ?, operation_order = ?, start_time = ?, finish_time = ?,
            duration_sec = ?, duration_min = ?, status = ?
        WHEN NOT MATCHED THEN INSERT (
            cycle_no, harness_code, operator_code, operation_name, operation_order,
            start_time, finish_time, duration_sec, duration_min, status
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """
        cycle_no = self.sql_int(row.get("cycle_no"))
        for name, order, start_key, finish_key, duration_key in operations:
            duration = self.sql_float(row.get(duration_key))
            if duration is None and not row.get(start_key) and not row.get(finish_key):
                continue
            status = "OK" if duration is not None and duration > 0 else "PENDING"
            params = (
                cycle_no,
                name,
                self.sql_text(row.get("harness_code"), 120),
                self.sql_text(row.get("operator"), 80),
                order,
                self.sql_dt(row.get(start_key)),
                self.sql_dt(row.get(finish_key)),
                duration,
                round(duration / 60.0, 3) if duration is not None else None,
                status,
                cycle_no,
                self.sql_text(row.get("harness_code"), 120),
                self.sql_text(row.get("operator"), 80),
                name,
                order,
                self.sql_dt(row.get(start_key)),
                self.sql_dt(row.get(finish_key)),
                duration,
                round(duration / 60.0, 3) if duration is not None else None,
                status,
            )
            self.sql_execute(query, params)

    def sql_write_maintenance(self, row: dict[str, Any]) -> None:
        query = """
        INSERT INTO dbo.Fact_Maintenance
            (event_time, cycle_no, harness_code, source, fault_code, fault_text, action_status)
        VALUES (?, ?, ?, ?, ?, ?, ?);
        """
        self.sql_execute(
            query,
            (
                self.sql_dt(row.get("scan_finish_time") or row.get("audit_finish_time") or now_iso()),
                self.sql_int(row.get("cycle_no")),
                self.sql_text(row.get("harness_code"), 120),
                self.sql_text(row.get("fault_source") or row.get("source"), 80),
                self.sql_int(row.get("fault_code")),
                self.sql_text(row.get("fault_text"), 500),
                "OPEN",
            ),
        )

    def sync_sql_from_csv(self, reason: str = "manual") -> None:
        if not self.cfg.sql_enabled:
            return
        if not self.ensure_sql_available(force=(reason == "startup")):
            return
        rows = self.load_cycles_from_csv()
        if not rows:
            return
        max_sql_cycle = self.sql_int(
            self.sql_scalar("SELECT ISNULL(MAX(cycle_no), 0) FROM dbo.Fact_Cycles")
        ) or 0
        written = 0
        for row in rows:
            cycle_no = self.sql_int(row.get("cycle_no"))
            if cycle_no is None or cycle_no <= max_sql_cycle:
                continue
            self.sql_write_cycle(row)
            if not self.sql_available:
                break
            written += 1
        if written:
            print(f"SQL sync from CSV ({reason}): {written} cycle(s) copied to SQL", flush=True)

    def sync_sql_from_csv_if_needed(self) -> None:
        if not self.cfg.sql_enabled:
            return
        now_mono = time.monotonic()
        if now_mono < self.sql_sync_at:
            return
        self.sql_sync_at = now_mono + 3.0
        try:
            mtime = self.cfg.csv_path.stat().st_mtime if self.cfg.csv_path.exists() else 0.0
        except Exception:
            mtime = 0.0
        if mtime <= self.last_sql_csv_mtime and self.sql_available:
            return
        self.last_sql_csv_mtime = mtime
        self.sync_sql_from_csv("csv_change")

    def begin_current_cycle(self) -> dict[str, Any]:
        plan = reference_clip_plan(self.state.current_code)
        if self.state.current_cycle and self.state.current_cycle.get("harness_code") == self.state.current_code:
            cycle = self.state.current_cycle
            cycle.setdefault("reference_family", reference_family(self.state.current_code))
            cycle.setdefault("required_clips", plan)
            cycle.setdefault("required_clips_json", json.dumps(plan, ensure_ascii=False))
            cycle.setdefault("clip_done", {})
            cycle.setdefault("clip_total_required", sum(plan.values()))
            cycle.setdefault("clip_total_done", 0)
            cycle.setdefault("clip_progress_text", clip_progress_text(cycle.get("clip_done", {}), plan))
            return cycle

        self.state.current_cycle = {
            "harness_code": self.state.current_code,
            "operator": self.state.current_operator,
            "comment": self.state.current_comment,
            "sequence_errors": self.state.sequence_errors,
            "reference_family": reference_family(self.state.current_code),
            "required_clips": plan,
            "required_clips_json": json.dumps(plan, ensure_ascii=False),
            "clip_done": {},
            "clip_total_required": sum(plan.values()),
            "clip_total_done": 0,
            "clip_progress_text": clip_progress_text({}, plan),
        }
        return self.state.current_cycle

    def topic(self, suffix: str) -> str:
        return f"{self.cfg.topic_prefix}/{suffix.strip('/')}"

    def ensure_csv(self) -> None:
        if not self.cfg.csv_path.exists():
            self.write_csv_rows([])
            return

        text = self.cfg.csv_path.read_text(encoding="utf-8-sig", errors="replace")
        if not text.strip():
            self.write_csv_rows([])
            return

        lines = text.splitlines()
        first_line = lines[0].strip().lower() if lines else ""
        data_lines = lines[1:] if first_line.startswith("sep=") else lines
        current_header = data_lines[0].lstrip("\ufeff").strip() if data_lines else ""
        expected_header = CSV_DELIMITER.join(CYCLE_COLUMNS)

        if first_line != EXCEL_SEPARATOR_LINE or current_header != expected_header:
            self.write_csv_rows(self.parse_csv_text(text))

    def parse_csv_text(self, text: str) -> list[dict[str, Any]]:
        lines = text.splitlines()
        if lines and lines[0].strip().lower().startswith("sep="):
            lines = lines[1:]
        if not lines:
            return []

        reader = csv.DictReader(lines, delimiter=CSV_DELIMITER)
        rows: list[dict[str, Any]] = []
        for row in reader:
            if not row:
                continue
            clean = {key: (row.get(key, "") or "") for key in CYCLE_COLUMNS}
            if any(str(value).strip() for value in clean.values()):
                rows.append(clean)
        return rows

    def load_cycles_from_csv(self) -> list[dict[str, Any]]:
        if not self.cfg.csv_path.exists():
            return []
        text = self.cfg.csv_path.read_text(encoding="utf-8-sig", errors="replace")
        return self.parse_csv_text(text)[-500:]

    def write_csv_rows(self, rows: list[dict[str, Any]]) -> None:
        with self.cfg.csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
            handle.write(EXCEL_SEPARATOR_LINE + "\r\n")
            writer = csv.DictWriter(handle, fieldnames=CYCLE_COLUMNS, delimiter=CSV_DELIMITER)
            writer.writeheader()
            for row in rows:
                writer.writerow({key: row.get(key, "") for key in CYCLE_COLUMNS})

    def rewrite_logs(self) -> None:
        self.write_csv_rows(self.state.cycles)

        if Workbook and load_workbook:
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "SEWS cycles"
                ws.append(CYCLE_COLUMNS)
                for row in self.state.cycles:
                    ws.append([row.get(key, "") for key in CYCLE_COLUMNS])
                try:
                    wb.save(self.cfg.xlsx_path)
                except PermissionError:
                    fallback = self.cfg.xlsx_path.with_name(self.cfg.xlsx_path.stem + "_latest.xlsx")
                    wb.save(fallback)
                    print(f"WARNING: XLSX locked by another app, wrote fallback: {fallback}", flush=True)
            except Exception as exc:
                print(f"WARNING: XLSX log skipped, CSV/web log is still OK: {exc}", flush=True)

    def connect_s7(self) -> bool:
        with self.s7_lock:
            if self.client.get_connected():
                self.state.s7_connected = True
                return True
            try:
                self.client.connect(self.cfg.s7_host, self.cfg.s7_rack, self.cfg.s7_slot)
                self.state.s7_connected = self.client.get_connected()
                self.state.last_s7_error = ""
                return self.state.s7_connected
            except Exception as exc:
                self.state.s7_connected = False
                self.state.last_s7_error = str(exc)
                return False

    def disconnect_s7(self) -> None:
        with self.s7_lock:
            try:
                self.client.disconnect()
            except Exception:
                pass
            self.state.s7_connected = False

    def write_db(self, db_number: int, offset: int, data: bytearray) -> bool:
        if not self.state.s7_connected:
            return False
        with self.s7_lock:
            # Keep Node-RED simulation responsive when PLCSIM Advanced is offline.
            # The main polling loop owns reconnect attempts; command writes only use
            # an already-established S7 connection.
            if not self.state.s7_connected:
                return False
            try:
                self.client.db_write(db_number, offset, data)
                self.state.s7_connected = True
                self.state.last_s7_error = ""
                return True
            except Exception as exc:
                self.state.s7_connected = False
                self.state.last_s7_error = str(exc)
                try:
                    self.client.disconnect()
                except Exception:
                    pass
                print(f"WARNING: S7 write failed DB{db_number} offset {offset}: {exc}", flush=True)
                return False

    def write_cmd_byte(self, index: int, value: bool) -> bool:
        return self.write_db(self.cfg.db_number, index, bytearray([1 if value else 0]))

    def mirror_to_sim_panel(self, db100_data: bytes, status: dict[str, Any]) -> None:
        # DB_SEWS_SIM_PANEL is now owned by the PLC FC so its layout can evolve safely.
        # Writing it from Python with hard-coded offsets would corrupt the easy TIA panel.
        return
        try:
            panel = bytearray(self.client.db_read(SIM_PANEL_DB_NUMBER, 0, SIM_PANEL_DB_SIZE))
            commands = {name: bool(db100_data[index]) for name, index in CMD_INDEX.items()}
            status_bool = status.get("bool", {})

            command_bits = {
                "enable": (0, 0),
                "auto_mode": (0, 1),
                "enable_datalog": (0, 2),
                "reset": (0, 3),
                "harness_present": (0, 5),
                "scan_start": (0, 6),
                "st2_done": (0, 7),
                "st3_done": (1, 0),
                "machine_ok": (1, 1),
                "scan_finish": (1, 2),
                "takeoff": (1, 3),
                "force_maintenance_fault": (1, 4),
                "force_machine_nok": (1, 5),
                "force_estop": (1, 6),
            }
            for name, (byte_index, bit_index) in command_bits.items():
                set_s7_bool(panel, byte_index, bit_index, commands.get(name, False))

            step = int(status.get("step", 0) or 0)
            status_bits = [
                ((1, 7), bool(status_bool.get("harness_present"))),
                ((2, 0), bool(status_bool.get("cycle_active"))),
                ((2, 1), step >= 20),
                ((2, 2), bool(status_bool.get("st2_working"))),
                ((2, 3), bool(status_bool.get("st3_working"))),
                ((2, 4), bool(status_bool.get("st4_testing"))),
                ((2, 5), bool(status_bool.get("machine_ok"))),
                ((2, 6), step >= 50),
                ((2, 7), bool(status_bool.get("st5_takeoff_ready"))),
                ((3, 0), bool(status_bool.get("maintenance_alarm"))),
                ((3, 1), bool(status_bool.get("fault_active"))),
            ]
            for (byte_index, bit_index), value in status_bits:
                set_s7_bool(panel, byte_index, bit_index, value)

            struct.pack_into(">h", panel, 4, step)
            struct.pack_into(">h", panel, 6, int(status.get("fault_code", 0) or 0))
            struct.pack_into(">i", panel, 8, int(status.get("cycle_count", 0) or 0))
            struct.pack_into(">i", panel, 12, int(status.get("cycle_timer_sec", 0) or 0))
            struct.pack_into(">i", panel, 16, int(status.get("st2_timer_sec", 0) or 0))
            struct.pack_into(">i", panel, 20, int(status.get("st3_timer_sec", 0) or 0))
            struct.pack_into(">i", panel, 24, int(status.get("machine_test_timer_sec", 0) or 0))
            struct.pack_into(">i", panel, 28, int(status.get("maintenance_timer_sec", 0) or 0))
            struct.pack_into(">i", panel, 32, int(status.get("last_duration_sec", 0) or 0))
            self.write_db(SIM_PANEL_DB_NUMBER, 0, panel)
        except Exception as exc:
            print(f"WARNING: could not mirror DB100 to DB_SEWS_SIM_PANEL DB{SIM_PANEL_DB_NUMBER}: {exc}", flush=True)

    def pulse_cmd(self, index: int, pulse_ms: int = 250) -> None:
        if not self.write_cmd_byte(index, True):
            return
        with self.state.lock:
            self.state.pending_pulses[index] = time.monotonic() + (pulse_ms / 1000.0)

    def clear_due_pulses(self) -> None:
        due: list[int] = []
        now = time.monotonic()
        with self.state.lock:
            for index, clear_at in list(self.state.pending_pulses.items()):
                if now >= clear_at:
                    due.append(index)
                    del self.state.pending_pulses[index]
        for index in due:
            self.write_cmd_byte(index, False)

    def initialize_plc_command_area(self) -> None:
        if not self.connect_s7():
            return
        cmd = bytearray(32)
        cmd[CMD_INDEX["enable"]] = 1
        cmd[CMD_INDEX["auto_mode"]] = 0
        cmd[CMD_INDEX["enable_datalog"]] = 1
        cmd[CMD_INDEX["bridge_active"]] = 1
        self.write_db(self.cfg.db_number, 0, cmd)
        self.state.pending_pulses.clear()

    def restore_normal_state(self) -> None:
        """Return the simulated line to a clean ready state after any fault."""
        if not self.connect_s7():
            return

        cmd = bytearray(32)
        cmd[CMD_INDEX["enable"]] = 1
        cmd[CMD_INDEX["auto_mode"]] = 0
        cmd[CMD_INDEX["enable_datalog"]] = 1
        cmd[CMD_INDEX["bridge_active"]] = 1
        self.write_db(self.cfg.db_number, 0, cmd)

        self.state.pending_pulses.clear()
        self.state.current_cycle = {}
        self.state.sequence_errors = 0
        self.state.pending_web_commit_at = 0.0
        self.state.pending_web_cycle_marker = -1

        # One reset pulse after all fault causes are removed.
        self.pulse_cmd(CMD_INDEX["reset"], 500)
        self.publish_bridge_status()

    def decode_status(self, data: bytes) -> dict[str, Any]:
        status_bool = {name: bool(data[32 + i]) for i, name in enumerate(STATUS_BOOL_NAMES)}
        status_bool["electric_test_running"] = status_bool.get("st4_testing", False)
        status_bool["electric_test_ok"] = status_bool.get("machine_ok", False)
        status_bool["electric_test_nok"] = status_bool.get("machine_nok", False)
        status_bool["clip_table_present"] = status_bool.get("clip_table_harness_present", False)
        status_bool["electrical_table_present"] = status_bool.get("electrical_table_harness_present", False)
        status_bool["ports_plugged"] = status_bool.get("electrical_ports_plugged", False)
        status_bool["prebloc_present"] = status_bool.get("prebloc_cable_present", False)
        status_bool["godet_cable_present"] = status_bool.get("prebloc_cable_present", False)
        status_bool["godet_sensor_ok"] = status_bool.get("prebloc_cable_present", False)
        status_bool["godet_cable_taken"] = status_bool.get("prebloc_ok", False)
        status_bool["conveyor_present"] = status_bool.get("conveyor_table_present", False)
        status_bool["audit_present"] = status_bool.get("audit_station_present", False)
        status_bool["pack_done"] = status_bool.get("pack_ok", False)
        status_int = [read_int(data, 64 + i * 2) for i in range(4)]
        status_dint = [read_dint(data, 128 + i * 4) for i in range(8)]
        status_real = [read_real(data, 256 + i * 4) for i in range(4)]
        commands = {name: bool(data[index]) for name, index in CMD_INDEX.items()}
        step = status_int[0]
        fault_code = status_int[2]
        commands = {name: bool(data[index]) for name, index in CMD_INDEX.items()}
        fault_source = "OK"
        if commands.get("force_estop") or status_bool.get("estop_active"):
            fault_source = "E-STOP"
            fault_text = "E-STOP ACTIVE - SAFETY STOP"
        elif commands.get("force_maintenance_fault") or status_bool.get("maintenance_alarm"):
            fault_source = "CAPTEUR"
            fault_text = "SENSOR FAULT - CHECK HARNESS/GODET/TABLE DETECTION"
        elif commands.get("force_machine_nok") or status_bool.get("machine_nok"):
            fault_source = "TEST ELEC"
            fault_text = "ELECTRICAL TEST NOK - CHECK ADAPTERS AND CURRENT"
        elif status_bool.get("fault_active") or fault_code:
            fault_source = "PLC FAULT"
            fault_text = STEP_TEXT.get(fault_code, f"Fault {fault_code}")
        else:
            fault_text = "NO FAULT"
        return {
            "ts": now_iso(),
            "connected": True,
            "s7_host": self.cfg.s7_host,
            "commands": commands,
            "bool": status_bool,
            "step": step,
            "step_text": STEP_TEXT.get(step, f"Step {step}"),
            "status_code": status_int[1],
            "fault_code": fault_code,
            "fault_text": fault_text,
            "fault_source": fault_source,
            "cycle_count": status_dint[0],
            "last_duration_sec": status_dint[1],
            "event_id": status_dint[2],
            "cycle_timer_sec": status_dint[3],
            "st2_timer_sec": status_dint[4],
            "st3_timer_sec": status_dint[5],
            "machine_test_timer_sec": status_dint[6],
            "electric_test_timer_sec": status_dint[6],
            "maintenance_timer_sec": status_dint[7],
            "last_duration_min": round(float(status_real[0]), 3),
            "harness_code": self.state.current_code,
        }

    def commit_cycle(self, status: dict[str, Any], cycle_no: int, source: str) -> None:
        self.state.current_cycle["cycle_no"] = cycle_no
        row = dict(self.state.current_cycle)
        row["cycle_no"] = cycle_no
        row.setdefault("harness_code", self.state.current_code)
        row.setdefault("operator", self.state.current_operator)
        row.setdefault("scan_finish_time", now_iso())
        required = row.get("required_clips")
        if not isinstance(required, dict):
            required = reference_clip_plan(row.get("harness_code", self.state.current_code))
        done = row.get("clip_done")
        if not isinstance(done, dict):
            done = {}
        row["reference_family"] = row.get("reference_family") or reference_family(row.get("harness_code", self.state.current_code))
        row["required_clips_json"] = json.dumps(required, ensure_ascii=False)
        row["clip_total_required"] = sum(int(v or 0) for v in required.values())
        row["clip_total_done"] = sum(int(done.get(k, 0) or 0) for k in required)
        row["clip_progress_text"] = clip_progress_text(done, required)
        row["plc_duration_sec"] = int(status.get("last_duration_sec", 0))
        prebloc_start = row.get("prebloc_start_epoch")
        prebloc_finish = row.get("prebloc_done_epoch")
        prebloc_duration = max(0.0, prebloc_finish - prebloc_start) if prebloc_start and prebloc_finish else 0.0
        conveyor_start = row.get("conveyor_start_epoch")
        conveyor_finish = row.get("conveyor_done_epoch")
        conveyor_duration = max(0.0, conveyor_finish - conveyor_start) if conveyor_start and conveyor_finish else 0.0
        tape_start = row.get("tape_start_epoch")
        tape_finish = row.get("tape_done_epoch")
        tape_duration = max(0.0, tape_finish - tape_start) if tape_start and tape_finish else 0.0
        clip_start = row.get("clip_first_clip_epoch")
        clip_finish = row.get("clip_scan_finish_epoch")
        clip_duration = max(0.0, clip_finish - clip_start) if clip_start and clip_finish else 0.0
        electric_start = row.get("electric_work_start_epoch")
        electric_finish = row.get("electric_scan_finish_epoch") or row.get("scan_finish_epoch")
        electric_duration = max(0.0, electric_finish - electric_start) if electric_start and electric_finish else 0.0
        audit_start = row.get("audit_start_epoch")
        audit_finish = row.get("audit_finish_epoch")
        audit_duration = max(0.0, audit_finish - audit_start) if audit_start and audit_finish else 0.0
        duration = prebloc_duration + conveyor_duration + tape_duration + clip_duration + electric_duration + audit_duration
        if duration <= 0.0:
            start = row.get("scan_start_epoch")
            finish = row.get("scan_finish_epoch") or time.time()
            duration = max(0.0, finish - start) if start else float(row["plc_duration_sec"])
        row["prebloc_operation_sec"] = round(prebloc_duration, 1) if prebloc_duration else ""
        row["prebloc_operation_min"] = round(prebloc_duration / 60.0, 3) if prebloc_duration else ""
        row["conveyor_operation_sec"] = round(conveyor_duration, 1) if conveyor_duration else ""
        row["conveyor_operation_min"] = round(conveyor_duration / 60.0, 3) if conveyor_duration else ""
        row["tape_operation_sec"] = round(tape_duration, 1) if tape_duration else ""
        row["tape_operation_min"] = round(tape_duration / 60.0, 3) if tape_duration else ""
        row["clip_operation_sec"] = round(clip_duration, 1) if clip_duration else ""
        row["clip_operation_min"] = round(clip_duration / 60.0, 3) if clip_duration else ""
        row["electric_operation_sec"] = round(electric_duration, 1) if electric_duration else ""
        row["electric_operation_min"] = round(electric_duration / 60.0, 3) if electric_duration else ""
        row["audit_operation_sec"] = round(audit_duration, 1) if audit_duration else ""
        row["audit_operation_min"] = round(audit_duration / 60.0, 3) if audit_duration else ""
        row["duration_work_sec"] = round(duration, 1)
        row["duration_work_min"] = round(duration / 60.0, 3)
        status_fault_code = int(status.get("fault_code", 0) or 0)
        row["fault_code"] = status_fault_code or int(row.get("fault_code") or 0)
        row["fault_text"] = status.get("fault_text", "") if status_fault_code else row.get("fault_text", "")
        row["fault_source"] = status.get("fault_source", "") if row["fault_code"] else row.get("fault_source", "")
        row["result"] = row.get("audit_result") or ("NOK" if status.get("fault_code", 0) else "OK")
        row["source"] = source
        row["sequence_errors"] = int(row.get("sequence_errors") or self.state.sequence_errors or 0)
        row.setdefault("comment", self.state.current_comment)
        for key in list(row):
            if key.endswith("_epoch"):
                del row[key]
        row.pop("required_clips", None)
        row.pop("clip_done", None)

        existing = next((i for i, item in enumerate(self.state.cycles) if item.get("cycle_no") == cycle_no), -1)
        if existing >= 0:
            self.state.cycles[existing] = row
        else:
            self.state.cycles.append(row)
        self.state.cycles = self.state.cycles[-500:]
        self.rewrite_logs()
        self.sql_write_cycle(row)
        self.publish_json("plc/cycle", row, retain=False)
        self.publish_history()
        self.state.pending_web_commit_at = 0.0
        self.state.sequence_errors = 0

    def update_cycle_from_status(self, status: dict[str, Any]) -> None:
        cycle_count = int(status.get("cycle_count", 0))
        if self.state.last_cycle_count < 0:
            self.state.last_cycle_count = cycle_count
            return
        if cycle_count <= self.state.last_cycle_count:
            return

        self.state.last_cycle_count = cycle_count
        self.commit_cycle(status, cycle_count, "plc_confirmed")

    def commit_web_fallback_if_due(self, status: dict[str, Any]) -> None:
        due_at = self.state.pending_web_commit_at
        if due_at <= 0.0 or time.monotonic() < due_at:
            return
        cycle_count = int(status.get("cycle_count", 0)) if self.state.s7_connected else 0
        if self.state.s7_connected and cycle_count > self.state.pending_web_cycle_marker:
            return
        existing_cycle_no = int(self.state.current_cycle.get("cycle_no") or 0)
        if existing_cycle_no > 0:
            cycle_no = existing_cycle_no
        else:
            cycle_no = max(self.state.last_cycle_count, cycle_count, len(self.state.cycles)) + 1
        self.state.last_cycle_count = cycle_no
        self.commit_cycle(status, cycle_no, "web_fallback")

    def publish_json(self, suffix: str, payload: dict[str, Any], retain: bool = False) -> None:
        self.mqtt.publish(self.topic(suffix), json.dumps(payload, ensure_ascii=False), retain=retain)

    def publish_bridge_status(self) -> None:
        payload = {
            "ts": now_iso(),
            "s7_connected": self.state.s7_connected,
            "s7_host": self.cfg.s7_host,
            "s7_db": self.cfg.db_number,
            "mqtt": f"{self.cfg.mqtt_host}:{self.cfg.mqtt_port}",
            "csv_path": str(self.cfg.csv_path),
            "xlsx_path": str(self.cfg.xlsx_path),
            "sql_enabled": self.cfg.sql_enabled,
            "sql_available": self.sql_available,
            "sql_last_error": self.sql_last_error,
            "current_code": self.state.current_code,
            "current_operator": self.state.current_operator,
            "current_comment": self.state.current_comment,
            "current_takt_target_sec": self.state.current_takt_target_sec,
            "sequence_errors": self.state.sequence_errors,
            "current_cycle": self.state.current_cycle,
            "last_error": self.state.last_s7_error,
            "cycles_logged": len(self.state.cycles),
        }
        self.publish_json("bridge/status", payload, retain=True)

    def publish_history(self) -> None:
        self.publish_json(
            "plc/history",
            {
                "ts": now_iso(),
                "cycles": self.state.cycles[-200:],
            },
            retain=True,
        )

    def on_mqtt_connect(self, client: mqtt.Client, userdata: Any, flags: Any, reason_code: Any, properties: Any) -> None:
        print(f"MQTT connected: {reason_code}")
        client.subscribe(self.topic("cmd/#"))
        self.publish_bridge_status()
        self.publish_history()

    def on_mqtt_message(self, client: mqtt.Client, userdata: Any, msg: mqtt.MQTTMessage) -> None:
        try:
            self._handle_mqtt_message(client, userdata, msg)
        except Exception as exc:
            self.state.last_s7_error = f"MQTT command error: {exc}"
            print(f"ERROR: MQTT command failed but bridge continues: {exc}", flush=True)
            try:
                self.publish_bridge_status()
            except Exception:
                pass

    def _handle_mqtt_message(self, client: mqtt.Client, userdata: Any, msg: mqtt.MQTTMessage) -> None:
        command = msg.topic.split("/")[-1]
        command = COMMAND_ALIASES.get(command, command)
        raw = msg.payload.decode("utf-8", errors="replace")
        try:
            payload: Any = json.loads(raw) if raw else {}
        except json.JSONDecodeError:
            payload = raw

        if command == "set_code":
            value = payload.get("value", payload) if isinstance(payload, dict) else payload
            new_code = str(value or self.cfg.default_harness_code).strip()
            previous_code = str(self.state.current_cycle.get("harness_code") or "")
            if self.state.current_cycle.get("cycle_no") or (previous_code and previous_code != new_code):
                self.state.current_cycle = {}
                self.state.pending_web_commit_at = 0.0
                self.state.pending_web_cycle_marker = -1
            self.state.current_code = new_code
            self.state.current_cycle["harness_code"] = self.state.current_code
            plan = reference_clip_plan(self.state.current_code)
            self.state.current_cycle["reference_family"] = reference_family(self.state.current_code)
            self.state.current_cycle["required_clips"] = plan
            self.state.current_cycle["required_clips_json"] = json.dumps(plan, ensure_ascii=False)
            self.state.current_cycle["clip_total_required"] = sum(plan.values())
            self.state.current_cycle.setdefault("clip_done", {})
            self.state.current_cycle["clip_total_done"] = sum(int(self.state.current_cycle["clip_done"].get(k, 0) or 0) for k in plan)
            self.state.current_cycle["clip_progress_text"] = clip_progress_text(self.state.current_cycle["clip_done"], plan)
            self.publish_bridge_status()
            return

        if command == "set_operator":
            value = payload.get("value", payload) if isinstance(payload, dict) else payload
            self.state.current_operator = str(value or "").strip()
            self.state.current_cycle["operator"] = self.state.current_operator
            self.publish_bridge_status()
            return

        if command == "set_comment":
            value = payload.get("value", payload) if isinstance(payload, dict) else payload
            self.state.current_comment = str(value or "").strip()
            self.state.current_cycle["comment"] = self.state.current_comment
            self.publish_bridge_status()
            return

        if command == "set_takt_target":
            value = payload.get("value", payload) if isinstance(payload, dict) else payload
            try:
                self.state.current_takt_target_sec = max(1.0, float(value))
            except Exception:
                self.state.current_takt_target_sec = 60.0
            self.state.current_cycle["takt_target_sec"] = self.state.current_takt_target_sec
            self.publish_bridge_status()
            return

        if command == "sequence_error":
            self.state.sequence_errors += 1
            self.state.current_cycle["sequence_errors"] = self.state.sequence_errors
            self.publish_bridge_status()
            return

        if command == "clip_done":
            clip = str((payload.get("clip") if isinstance(payload, dict) else payload) or "").upper().strip()
            if clip not in {f"C{i}" for i in range(1, 12)}:
                print(f"Unknown clip ignored: {clip}")
                return
            plan = self.state.current_cycle.get("required_clips")
            if not isinstance(plan, dict):
                plan = reference_clip_plan(self.state.current_code)
                self.state.current_cycle["required_clips"] = plan
            if int(plan.get(clip, 0) or 0) <= 0:
                self.publish_bridge_status()
                return
            done = self.state.current_cycle.setdefault("clip_done", {})
            done[clip] = min(int(plan[clip]), int(done.get(clip, 0) or 0) + 1)
            event_time = now_iso()
            event_epoch = time.time()
            first_clip = "clip_first_clip_time" not in self.state.current_cycle
            self.state.current_cycle.setdefault("clip_first_clip_time", event_time)
            self.state.current_cycle.setdefault("clip_first_clip_epoch", event_epoch)
            if first_clip:
                self.state.current_cycle.setdefault("st2_done_time", event_time)
                self.pulse_cmd(CMD_INDEX["st2_done"], DEFAULT_PULSE_MS["st2_done"])
            self.state.current_cycle["clip_total_required"] = sum(int(v or 0) for v in plan.values())
            self.state.current_cycle["clip_total_done"] = sum(int(done.get(k, 0) or 0) for k in plan)
            self.state.current_cycle["clip_progress_text"] = clip_progress_text(done, plan)
            if (
                self.state.current_cycle["clip_total_required"] > 0
                and self.state.current_cycle["clip_total_done"] >= self.state.current_cycle["clip_total_required"]
                and "st3_done_time" not in self.state.current_cycle
            ):
                self.state.current_cycle["st3_done_time"] = event_time
                self.pulse_cmd(CMD_INDEX["st3_done"], DEFAULT_PULSE_MS["st3_done"])
            self.publish_bridge_status()
            return

        if command in {"audit2_ok", "audit3_ok", "pack_ok"}:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            field_map = {
                "audit2_ok": ("audit2_finish_time", "audit2_finish_epoch"),
                "audit3_ok": ("audit3_finish_time", "audit3_finish_epoch"),
                "pack_ok": ("pack_time", "pack_epoch"),
            }
            time_field, epoch_field = field_map[command]
            cycle[time_field] = event_time
            cycle[epoch_field] = event_epoch
            if command == "pack_ok":
                cycle["pack_result"] = "OK"
                cycle.setdefault("result", "OK")
                self.state.pending_web_commit_at = time.monotonic() + 1.0
                self.state.pending_web_cycle_marker = self.state.last_cycle_count
            current_cycle_no = str(cycle.get("cycle_no") or "")
            if (
                current_cycle_no
                and self.state.cycles
                and str(self.state.cycles[-1].get("cycle_no") or "") == current_cycle_no
            ):
                self.state.cycles[-1][time_field] = event_time
                if command == "pack_ok":
                    self.state.cycles[-1]["pack_result"] = "OK"
                    self.state.cycles[-1].setdefault("result", "OK")
                self.rewrite_logs()
                self.publish_json("plc/cycle", self.state.cycles[-1], retain=False)
                self.publish_history()
            self.pulse_cmd(CMD_INDEX[command], DEFAULT_PULSE_MS[command])
            self.publish_bridge_status()
            return

        if command == "clear_faults":
            for name in ("force_maintenance_fault", "force_machine_nok", "force_estop"):
                self.write_cmd_byte(CMD_INDEX[name], False)
            self.write_cmd_byte(CMD_INDEX["machine_ok"], False)
            self.pulse_cmd(CMD_INDEX["reset"], 250)
            return

        if command == "restore_normal":
            self.restore_normal_state()
            return

        if command == "all_stop":
            for index in range(0, 32):
                self.write_cmd_byte(index, False)
            self.state.current_cycle = {}
            self.state.pending_web_commit_at = 0.0
            return

        if command not in CMD_INDEX:
            print(f"Unknown command ignored: {command}")
            return

        value = parse_bool_payload(payload, True)
        pulse_ms = DEFAULT_PULSE_MS.get(command, 250)
        if isinstance(payload, dict) and "pulse_ms" in payload:
            pulse_ms = int(payload["pulse_ms"])

        if command == "reset" and value:
            self.state.current_cycle = {}
            self.state.pending_web_commit_at = 0.0
            self.state.pending_web_cycle_marker = -1
        elif command == "prebloc_cable_present" and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["prebloc_start_time"] = event_time
            cycle["prebloc_start_epoch"] = event_epoch
            cycle["godet_loaded_time"] = event_time
            cycle["godet_sensor_detect_time"] = event_time
            cycle.setdefault("cycle_start_epoch", event_epoch)
        elif command == "prebloc_done" and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["prebloc_done_time"] = event_time
            cycle["prebloc_done_epoch"] = event_epoch
            cycle["godet_taken_time"] = event_time
            cycle["chariot_loaded_time"] = event_time
        elif command == "move_to_conveyor" and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["conveyor_start_time"] = event_time
            cycle["layup_start_time"] = event_time
            cycle["conveyor_start_epoch"] = event_epoch
            cycle["layup_start_epoch"] = event_epoch
        elif command == "conveyor_done" and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["conveyor_done_time"] = event_time
            cycle["layup_done_time"] = event_time
            cycle["conveyor_done_epoch"] = event_epoch
            cycle["layup_done_epoch"] = event_epoch
        elif command == "harness_present" and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["deposit_time"] = event_time
            cycle["clip_table_time"] = event_time
            cycle["deposit_epoch"] = event_epoch
        elif command == "scan_start":
            self.begin_current_cycle()
            self.state.current_cycle.setdefault("harness_code", self.state.current_code)
            event_time = now_iso()
            event_epoch = time.time()
            self.state.current_cycle["scan_start_time"] = event_time
            self.state.current_cycle["clip_scan_start_time"] = event_time
            self.state.current_cycle.setdefault("tape_start_time", event_time)
            self.state.current_cycle.setdefault("tape_start_epoch", event_epoch)
            self.state.current_cycle["scan_start_epoch"] = event_epoch
            self.state.current_cycle["clip_scan_start_epoch"] = event_epoch
        elif command == "st2_done":
            event_time = now_iso()
            event_epoch = time.time()
            self.state.current_cycle["st2_done_time"] = event_time
            self.state.current_cycle["tape_done_time"] = event_time
            self.state.current_cycle["tape_done_epoch"] = event_epoch
        elif command == "st3_done":
            event_time = now_iso()
            self.state.current_cycle["st3_done_time"] = event_time
            self.state.current_cycle["clip_done_time"] = event_time
        elif command == "clip_finish_scan":
            self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            self.state.current_cycle["clip_scan_finish_time"] = event_time
            self.state.current_cycle["clip_scan_finish_epoch"] = event_epoch
            self.state.current_cycle.setdefault("scan_clip_finish_time", event_time)
        elif command == "transfer_to_electrical" and value:
            self.begin_current_cycle()
            self.state.current_cycle["transfer_to_electrical_time"] = now_iso()
        elif command == "electrical_harness_present" and value:
            self.begin_current_cycle()
            self.state.current_cycle["electrical_table_time"] = now_iso()
        elif command == "electric_start_scan":
            self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            self.state.current_cycle["electric_scan_start_time"] = event_time
            self.state.current_cycle["electric_scan_start_epoch"] = event_epoch
        elif command == "plug_connectors" and value:
            self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            self.state.current_cycle["connectors_plugged_time"] = event_time
            self.state.current_cycle.setdefault("electric_work_start_time", event_time)
            self.state.current_cycle.setdefault("electric_work_start_epoch", event_epoch)
            status_bool = self.state.last_status.get("bool", {}) if isinstance(self.state.last_status, dict) else {}
            status_cmd = self.state.last_status.get("commands", {}) if isinstance(self.state.last_status, dict) else {}
            nok_active = bool(
                status_bool.get("fault_active")
                or status_bool.get("maintenance_alarm")
                or status_bool.get("machine_nok")
                or status_bool.get("estop_active")
                or status_cmd.get("force_machine_nok")
                or status_cmd.get("force_maintenance_fault")
                or status_cmd.get("force_estop")
            )
            if not nok_active:
                self.state.current_cycle.setdefault("electric_test_time", event_time)
                self.write_cmd_byte(CMD_INDEX["machine_ok"], True)
            else:
                self.write_cmd_byte(CMD_INDEX["machine_ok"], False)
        elif command == "machine_ok" and value:
            self.begin_current_cycle()
            self.state.current_cycle["electric_test_time"] = now_iso()
        elif command in {"scan_finish", "electric_finish_scan"}:
            self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            self.state.current_cycle["scan_finish_time"] = event_time
            self.state.current_cycle["electric_scan_finish_time"] = event_time
            self.state.current_cycle["scan_finish_epoch"] = event_epoch
            self.state.current_cycle["electric_scan_finish_epoch"] = event_epoch
        elif command == "audit_start" and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["audit_start_time"] = event_time
            cycle["audit1_start_time"] = event_time
            cycle["audit_start_epoch"] = event_epoch
            cycle["audit1_start_epoch"] = event_epoch
        elif command in {"audit_ok", "audit_nok"} and value:
            cycle = self.begin_current_cycle()
            event_time = now_iso()
            event_epoch = time.time()
            cycle["audit_finish_time"] = event_time
            cycle["audit1_finish_time"] = event_time
            cycle["audit_finish_epoch"] = event_epoch
            cycle["audit1_finish_epoch"] = event_epoch
            cycle["audit_result"] = "NOK" if command == "audit_nok" else "OK"
            if command == "audit_nok":
                cycle["result"] = "NOK"
                cycle["fault_text"] = "AUDIT NOK - QUALITY CONTROL FAILED"
                cycle["fault_code"] = 92
            self.state.pending_web_commit_at = time.monotonic() + 1.5
            self.state.pending_web_cycle_marker = self.state.last_cycle_count
        elif command in {"force_machine_nok", "force_maintenance_fault", "force_estop"} and value:
            if self.state.current_cycle:
                event_time = now_iso()
                self.state.current_cycle.setdefault("electric_test_time", event_time)
                self.state.current_cycle.setdefault("scan_finish_time", event_time)
                self.state.current_cycle.setdefault("electric_scan_finish_time", event_time)
                self.state.current_cycle["scan_finish_epoch"] = time.time()
                self.state.current_cycle["electric_scan_finish_epoch"] = self.state.current_cycle["scan_finish_epoch"]
                self.state.pending_web_commit_at = time.monotonic() + 1.8
                self.state.pending_web_cycle_marker = self.state.last_cycle_count
        elif command == "takeoff":
            cycle = self.begin_current_cycle()
            takeoff_time = now_iso()
            cycle["takeoff_time"] = takeoff_time
            cycle["takeoff_epoch"] = time.time()

            current_cycle_no = str(cycle.get("cycle_no") or "")
            target_index = -1
            if current_cycle_no:
                target_index = next(
                    (i for i, item in enumerate(self.state.cycles) if str(item.get("cycle_no") or "") == current_cycle_no),
                    -1,
                )
            if target_index < 0:
                current_code = str(cycle.get("harness_code") or self.state.current_code or "")
                for i in range(len(self.state.cycles) - 1, -1, -1):
                    item = self.state.cycles[i]
                    if str(item.get("harness_code") or "") == current_code:
                        target_index = i
                        break

            if target_index >= 0:
                self.state.cycles[target_index]["takeoff_time"] = takeoff_time
                self.rewrite_logs()
                self.sql_write_cycle(self.state.cycles[target_index])
                self.publish_json("plc/cycle", self.state.cycles[target_index], retain=False)
                self.publish_history()
            else:
                self.state.pending_web_commit_at = time.monotonic() + 0.8
                self.state.pending_web_cycle_marker = self.state.last_cycle_count

            self.publish_bridge_status()
            for name in ("harness_present", "electrical_harness_present", "machine_ok"):
                self.write_cmd_byte(CMD_INDEX[name], False)

        if command in PULSE_COMMANDS:
            self.pulse_cmd(CMD_INDEX[command], pulse_ms)
        else:
            self.write_cmd_byte(CMD_INDEX[command], value)

    def run(self) -> int:
        print("SEWS S7-1500 MQTT bridge")
        print(f"S7 target: {self.cfg.s7_host} rack {self.cfg.s7_rack} slot {self.cfg.s7_slot} DB{self.cfg.db_number}")
        print(f"MQTT: {self.cfg.mqtt_host}:{self.cfg.mqtt_port} prefix {self.cfg.topic_prefix}")
        print(f"CSV: {self.cfg.csv_path}")

        self.mqtt.connect(self.cfg.mqtt_host, self.cfg.mqtt_port, keepalive=30)
        self.mqtt.loop_start()
        self.initialize_plc_command_area()

        while self.state.running:
            try:
                self.clear_due_pulses()
                status = None
                with self.s7_lock:
                    if self.connect_s7():
                        data = bytes(self.client.db_read(self.cfg.db_number, 0, self.cfg.db_size))
                        status = self.decode_status(data)
                        self.mirror_to_sim_panel(data, status)
                if status is not None:
                    self.state.last_status = status
                    self.update_cycle_from_status(status)
                    self.commit_web_fallback_if_due(status)
                    self.publish_json("plc/status", status, retain=True)
                    self.sql_write_live_status(status)
                else:
                    offline_status = {
                        "ts": now_iso(),
                        "connected": False,
                        "s7_host": self.cfg.s7_host,
                        "error": self.state.last_s7_error,
                        "harness_code": self.state.current_code,
                    }
                    self.commit_web_fallback_if_due(offline_status)
                    self.publish_json(
                        "plc/status",
                        offline_status,
                        retain=True,
                    )
                    self.sql_write_live_status(offline_status)
                self.sync_sql_from_csv_if_needed()
                self.publish_bridge_status()
                time.sleep(self.cfg.poll_ms / 1000.0)
            except KeyboardInterrupt:
                break
            except Exception as exc:
                self.state.s7_connected = False
                self.state.last_s7_error = str(exc)
                print(f"Bridge loop error: {exc}", flush=True)
                self.disconnect_s7()
                time.sleep(1.0)

        self.mqtt.loop_stop()
        try:
            self.client.disconnect()
        except Exception:
            pass
        return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True, help="Path to config.s7_1500_plcsim_advanced.json")
    args = parser.parse_args()
    cfg = BridgeConfig.load(args.config)
    mutex_handle = acquire_named_mutex()
    if mutex_handle < 0:
        return 3
    lock_path = cfg.csv_path.parent / "sews_s7_1500_bridge.lock"
    lock_fd = acquire_lock(lock_path)
    if lock_fd < 0:
        release_named_mutex(mutex_handle)
        return 3
    bridge = SewsBridge(cfg)

    def stop_handler(signum: int, frame: Any) -> None:
        bridge.state.running = False

    signal.signal(signal.SIGINT, stop_handler)
    signal.signal(signal.SIGTERM, stop_handler)
    try:
        return bridge.run()
    finally:
        try:
            os.close(lock_fd)
        except Exception:
            pass
        try:
            lock_path.unlink()
        except Exception:
            pass
        release_named_mutex(mutex_handle)


if __name__ == "__main__":
    raise SystemExit(main())
